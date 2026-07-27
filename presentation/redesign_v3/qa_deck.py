from __future__ import annotations

import json
import re
import zipfile
from pathlib import Path
from typing import Any

from PIL import Image
from pypdf import PdfReader
from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE


ROOT = Path(__file__).resolve().parent
PPTX = ROOT / "DPS_Studio_组会汇报.pptx"
PDF = ROOT / "DPS_Studio_组会汇报.pdf"
RENDERS = ROOT / "rendered_slides"
OUTPUT = ROOT / "_qa_tmp" / "qa_summary.json"

BANNED = (
    "赋能",
    "闭环",
    "全链路",
    "底座",
    "抓手",
    "生态",
    "智能化",
    "系统化提升",
    "显著提升",
    "一站式",
    "高效可靠",
    "关键价值",
    "质量门",
)


def natural_key(path: Path) -> list[int | str]:
    return [
        int(part) if part.isdigit() else part.lower()
        for part in re.split(r"(\d+)", path.name)
    ]


def slide_text(slide: Any) -> str:
    parts: list[str] = []
    for shape in slide.shapes:
        if getattr(shape, "has_text_frame", False):
            for paragraph in shape.text_frame.paragraphs:
                if paragraph.text:
                    parts.append(paragraph.text)
        if getattr(shape, "has_table", False):
            for row in shape.table.rows:
                for cell in row.cells:
                    parts.append(cell.text)
    return "\n".join(parts)


def main() -> None:
    deck = Presentation(PPTX)
    slide_width = deck.slide_width
    slide_height = deck.slide_height
    texts = [slide_text(slide) for slide in deck.slides]
    all_text = "\n".join(texts)

    pictures_per_slide: list[int] = []
    native_per_slide: list[int] = []
    full_slide_picture_shapes: list[dict[str, int]] = []
    out_of_bounds: list[dict[str, int]] = []
    text_shape_count = 0
    picture_count = 0
    for slide_number, slide in enumerate(deck.slides, 1):
        pictures = 0
        native = 0
        for shape_number, shape in enumerate(slide.shapes, 1):
            if shape.shape_type == MSO_SHAPE_TYPE.PICTURE:
                pictures += 1
                picture_count += 1
                area_ratio = (shape.width * shape.height) / (
                    slide_width * slide_height
                )
                if area_ratio > 0.9:
                    full_slide_picture_shapes.append(
                        {
                            "slide": slide_number,
                            "shape": shape_number,
                            "area_ratio": round(area_ratio, 4),
                        }
                    )
            else:
                native += 1
            if getattr(shape, "has_text_frame", False):
                text_shape_count += 1
            if (
                shape.left < 0
                or shape.top < 0
                or shape.left + shape.width > slide_width + 1000
                or shape.top + shape.height > slide_height + 1000
            ):
                out_of_bounds.append(
                    {
                        "slide": slide_number,
                        "shape": shape_number,
                        "left": shape.left,
                        "top": shape.top,
                        "width": shape.width,
                        "height": shape.height,
                    }
                )
        pictures_per_slide.append(pictures)
        native_per_slide.append(native)

    rendered = sorted(RENDERS.glob("*.PNG"), key=natural_key)
    render_sizes: list[list[int]] = []
    for path in rendered:
        with Image.open(path) as image:
            render_sizes.append([image.width, image.height])

    with zipfile.ZipFile(PPTX) as archive:
        note_slides = [
            name
            for name in archive.namelist()
            if re.fullmatch(r"ppt/notesSlides/notesSlide\d+\.xml", name)
        ]

    banned_hits = {
        word: [index + 1 for index, text in enumerate(texts) if word in text]
        for word in BANNED
        if word in all_text
    }

    required_phrases = {
        "临时波长边界": "演示值" in all_text and "尚未确认" in all_text,
        "LiF 边界": "无 LiF 修正" in all_text or "不含 LiF" in all_text,
        "表观速度命名": "表观速度" in all_text,
        "双通道不融合": "不平均、不融合" in all_text,
        "dB 非 SNR": "不是 SNR" in all_text,
        "显示零与 NaN": "显示假设" in all_text and "NaN" in all_text,
    }

    payload = {
        "pptx": str(PPTX),
        "pdf": str(PDF),
        "slide_count": len(deck.slides),
        "pdf_page_count": len(PdfReader(PDF).pages),
        "render_count": len(rendered),
        "render_sizes": sorted({tuple(size) for size in render_sizes}),
        "picture_count": picture_count,
        "pictures_per_slide": pictures_per_slide,
        "native_objects_per_slide": native_per_slide,
        "text_shape_count": text_shape_count,
        "full_slide_picture_shapes": full_slide_picture_shapes,
        "out_of_bounds": out_of_bounds,
        "notes_slide_count": len(note_slides),
        "banned_hits": banned_hits,
        "required_phrase_checks": required_phrases,
        "all_required_phrase_checks_passed": all(required_phrases.values()),
        "pptx_bytes": PPTX.stat().st_size,
        "pdf_bytes": PDF.stat().st_size,
    }
    OUTPUT.write_text(
        json.dumps(payload, ensure_ascii=False, indent=2), encoding="utf-8"
    )
    print(json.dumps(payload, ensure_ascii=False, indent=2))

    if len(deck.slides) != 18:
        raise SystemExit("PPTX slide count is not 18")
    if len(PdfReader(PDF).pages) != 18:
        raise SystemExit("PDF page count is not 18")
    if len(rendered) != 18 or any(size != [1920, 1080] for size in render_sizes):
        raise SystemExit("Rendered slides are incomplete or have wrong dimensions")
    if full_slide_picture_shapes:
        raise SystemExit("Detected a full-slide raster image")
    if out_of_bounds:
        raise SystemExit("Detected out-of-bounds shapes")
    if banned_hits:
        raise SystemExit("Detected banned wording")
    if not all(required_phrases.values()):
        raise SystemExit("Missing required scientific boundary statement")
    if len(note_slides) != 18:
        raise SystemExit("Speaker notes are incomplete")


if __name__ == "__main__":
    main()
