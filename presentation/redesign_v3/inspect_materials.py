from __future__ import annotations

import csv
import hashlib
import json
import math
from collections import defaultdict
from pathlib import Path
from typing import Any

from PIL import Image, ImageDraw, ImageFont, ImageOps
from pptx import Presentation


REPO = Path(__file__).resolve().parents[2]
OUT = Path(__file__).resolve().parent
DATA = REPO / "presentations" / "data"
QA = OUT / "_qa_tmp"


def sha256(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as handle:
        for block in iter(lambda: handle.read(1024 * 1024), b""):
            digest.update(block)
    return digest.hexdigest()


def average_hash(path: Path, size: int = 16) -> str:
    with Image.open(path) as image:
        gray = image.convert("L").resize((size, size), Image.Resampling.LANCZOS)
        pixels = list(gray.getdata())
    threshold = sum(pixels) / len(pixels)
    bits = "".join("1" if value >= threshold else "0" for value in pixels)
    return f"{int(bits, 2):0{size * size // 4}x}"


def image_record(path: Path) -> dict[str, Any]:
    with Image.open(path) as image:
        width, height = image.size
        mode = image.mode
    return {
        "path": path.relative_to(REPO).as_posix(),
        "bytes": path.stat().st_size,
        "width": width,
        "height": height,
        "aspect_ratio": round(width / height, 4),
        "mode": mode,
        "sha256": sha256(path),
        "average_hash": average_hash(path),
    }


def parse_number(value: str) -> float | None:
    text = value.strip()
    if not text:
        return None
    try:
        number = float(text)
    except ValueError:
        return None
    return number if math.isfinite(number) else None


def csv_record(path: Path) -> dict[str, Any]:
    encodings = ("utf-8-sig", "utf-8", "gb18030")
    rows: list[list[str]] = []
    encoding_used = ""
    last_error: Exception | None = None
    for encoding in encodings:
        try:
            with path.open("r", encoding=encoding, newline="") as handle:
                rows = list(csv.reader(handle))
            encoding_used = encoding
            break
        except UnicodeDecodeError as error:
            last_error = error
    if not encoding_used:
        raise RuntimeError(f"Cannot decode {path}: {last_error}")

    header = rows[0] if rows else []
    data_rows = rows[1:] if rows else []
    numeric_summary: dict[str, dict[str, float | int | None]] = {}
    for index, column in enumerate(header):
        values = [
            number
            for row in data_rows
            if index < len(row)
            for number in [parse_number(row[index])]
            if number is not None
        ]
        if values:
            numeric_summary[column] = {
                "count": len(values),
                "min": min(values),
                "max": max(values),
                "mean": sum(values) / len(values),
            }

    return {
        "path": path.relative_to(REPO).as_posix(),
        "bytes": path.stat().st_size,
        "encoding": encoding_used,
        "row_count": len(data_rows),
        "column_count": len(header),
        "header": header,
        "numeric_summary": numeric_summary,
        "sha256": sha256(path),
    }


def json_record(path: Path) -> dict[str, Any]:
    payload = json.loads(path.read_text(encoding="utf-8"))
    return {
        "path": path.relative_to(REPO).as_posix(),
        "bytes": path.stat().st_size,
        "top_level_type": type(payload).__name__,
        "top_level_keys": sorted(payload) if isinstance(payload, dict) else [],
        "payload": payload,
        "sha256": sha256(path),
    }


def pptx_record(path: Path) -> dict[str, Any]:
    deck = Presentation(path)
    slides: list[dict[str, Any]] = []
    for number, slide in enumerate(deck.slides, 1):
        text_blocks: list[str] = []
        pictures = 0
        tables = 0
        charts = 0
        native_shapes = 0
        for shape in slide.shapes:
            if shape.shape_type == 13:
                pictures += 1
            else:
                native_shapes += 1
            if getattr(shape, "has_table", False):
                tables += 1
            if getattr(shape, "has_chart", False):
                charts += 1
            if getattr(shape, "has_text_frame", False):
                text = "\n".join(
                    paragraph.text for paragraph in shape.text_frame.paragraphs
                ).strip()
                if text:
                    text_blocks.append(text)
        slides.append(
            {
                "number": number,
                "pictures": pictures,
                "native_shapes": native_shapes,
                "tables": tables,
                "charts": charts,
                "text": text_blocks,
            }
        )
    return {
        "path": path.relative_to(REPO).as_posix(),
        "bytes": path.stat().st_size,
        "slide_count": len(deck.slides),
        "width": deck.slide_width,
        "height": deck.slide_height,
        "slides": slides,
        "sha256": sha256(path),
    }


def load_font(size: int) -> ImageFont.ImageFont:
    candidates = (
        Path("C:/Windows/Fonts/msyh.ttc"),
        Path("C:/Windows/Fonts/msyh.ttf"),
        Path("C:/Windows/Fonts/arial.ttf"),
    )
    for candidate in candidates:
        if candidate.exists():
            return ImageFont.truetype(str(candidate), size)
    return ImageFont.load_default()


def make_contact_sheets(records: list[dict[str, Any]]) -> list[str]:
    paths = [REPO / record["path"] for record in records]
    font = load_font(20)
    sheets: list[str] = []
    columns, rows = 3, 3
    page_w, page_h = 1800, 1350
    margin = 30
    cell_w = (page_w - margin * (columns + 1)) // columns
    cell_h = (page_h - margin * (rows + 1)) // rows
    image_h = cell_h - 74
    for page_index in range(0, len(paths), columns * rows):
        canvas = Image.new("RGB", (page_w, page_h), "#f4f0e8")
        draw = ImageDraw.Draw(canvas)
        for local_index, path in enumerate(paths[page_index : page_index + 9]):
            row = local_index // columns
            column = local_index % columns
            x = margin + column * (cell_w + margin)
            y = margin + row * (cell_h + margin)
            with Image.open(path) as source:
                preview = ImageOps.contain(
                    source.convert("RGB"),
                    (cell_w, image_h),
                    Image.Resampling.LANCZOS,
                )
            px = x + (cell_w - preview.width) // 2
            py = y + (image_h - preview.height) // 2
            canvas.paste(preview, (px, py))
            label = path.relative_to(DATA).as_posix()
            if len(label) > 72:
                label = "…" + label[-71:]
            draw.text((x, y + image_h + 8), label, fill="#222222", font=font)
        output = QA / f"data_contact_sheet_{page_index // 9 + 1:02d}.png"
        canvas.save(output, quality=92)
        sheets.append(output.relative_to(REPO).as_posix())
    return sheets


def main() -> None:
    QA.mkdir(parents=True, exist_ok=True)
    images = sorted(
        [
            image_record(path)
            for path in DATA.rglob("*")
            if path.is_file() and path.suffix.lower() in {".png", ".jpg", ".jpeg"}
        ],
        key=lambda item: item["path"],
    )
    csvs = sorted(
        [csv_record(path) for path in DATA.rglob("*.csv")],
        key=lambda item: item["path"],
    )
    jsons = sorted(
        [json_record(path) for path in DATA.rglob("*.json")],
        key=lambda item: item["path"],
    )

    extra_files = [
        REPO
        / "presentations"
        / "group_meeting"
        / "20260726_redesign"
        / "assets"
        / "real"
        / "raw_voltage_dual_channel.png",
        REPO / "data" / "reference" / "legacy" / "legacy_velocity_time.csv",
        REPO / "data" / "reference" / "legacy" / "README.md",
        REPO / "docs" / "旧软件速度时间数据.csv",
        REPO / "README.md",
        REPO / "AGENTS.md",
        REPO / "pyproject.toml",
        REPO / "configs" / "demo_dual_profile.toml",
    ]
    extras: list[dict[str, Any]] = []
    for path in extra_files:
        if not path.exists():
            extras.append(
                {"path": path.relative_to(REPO).as_posix(), "exists": False}
            )
            continue
        item: dict[str, Any] = {
            "path": path.relative_to(REPO).as_posix(),
            "exists": True,
            "bytes": path.stat().st_size,
            "sha256": sha256(path),
        }
        if path.suffix.lower() in {".png", ".jpg", ".jpeg"}:
            item["image"] = image_record(path)
        elif path.suffix.lower() == ".csv":
            item["csv"] = csv_record(path)
        else:
            text = path.read_text(encoding="utf-8-sig")
            item["line_count"] = len(text.splitlines())
            item["text"] = text
        extras.append(item)

    decks = [
        pptx_record(path)
        for path in (
            REPO / "presentations" / "Sigapore_presentation.pptx",
            REPO
            / "presentations"
            / "group_meeting"
            / "20260725"
            / "DPS_Studio_组会汇报_20260725.pptx",
            REPO
            / "presentations"
            / "group_meeting"
            / "20260726_redesign"
            / "DPS_Studio_组会汇报_redesign.pptx",
        )
        if path.exists()
    ]

    exact_duplicate_groups = [
        paths
        for paths in defaultdict(list).values()
        if len(paths) > 1
    ]
    by_sha: dict[str, list[str]] = defaultdict(list)
    by_ahash: dict[str, list[str]] = defaultdict(list)
    for record in images:
        by_sha[record["sha256"]].append(record["path"])
        by_ahash[record["average_hash"]].append(record["path"])
    exact_duplicate_groups = [paths for paths in by_sha.values() if len(paths) > 1]
    visual_duplicate_groups = [
        paths for paths in by_ahash.values() if len(paths) > 1
    ]

    payload = {
        "repository": str(REPO),
        "data_directory": DATA.relative_to(REPO).as_posix(),
        "counts": {
            "images": len(images),
            "csvs": len(csvs),
            "jsons": len(jsons),
        },
        "images": images,
        "csvs": csvs,
        "jsons": jsons,
        "extras": extras,
        "decks": decks,
        "exact_duplicate_groups": exact_duplicate_groups,
        "same_average_hash_groups": visual_duplicate_groups,
        "contact_sheets": make_contact_sheets(images),
    }
    output = QA / "material_inspection.json"
    output.write_text(
        json.dumps(payload, ensure_ascii=False, indent=2), encoding="utf-8"
    )
    print(
        json.dumps(
            {
                "output": str(output),
                "counts": payload["counts"],
                "pptx_slides": {
                    deck["path"]: deck["slide_count"] for deck in decks
                },
                "exact_duplicate_groups": len(exact_duplicate_groups),
                "same_average_hash_groups": len(visual_duplicate_groups),
                "contact_sheets": payload["contact_sheets"],
            },
            ensure_ascii=False,
            indent=2,
        )
    )


if __name__ == "__main__":
    main()
